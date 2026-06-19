from __future__ import annotations

import argparse
import json
import re
from copy import deepcopy
from html import unescape
from pathlib import Path
from typing import Any

from openpyxl.cell.cell import MergedCell
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries


HEADER_SHEET = "Datos de cotización"
SYSTEMS_INDEX_SHEET = "Índice de sistemas"
PAYLOAD_SHEET = "FDI_Payload_JSON"
LAYOUT_FILE = Path(__file__).with_name("fdi-excel-layout.json")

TYPE_TEMPLATES = {
    "SEL": "SEL_S01_Form",
    "DIN": "DIN_S01_Form",
    "PBK": "PBK_S01_Form",
    "DRV": "DRV_S01_Form",
    "CAN": "CAN_S01_Form",
    "MEZ": "MEZ_S01_Form",
    "CFL": "CFL_S01_Form",
    "MZL": "MZL_S01_Form",
}

HEADER_ROWS = [
    (6, ("Folio",)),
    (7, ("VendedoresLookUp", "Vendedor", "VendedorNombre")),
    (8, ("EmpresaLookUp", "Cliente", "Title")),
    (9, ("ContactoLookUp", "Nombre de contacto", "Contacto")),
    (10, ("Dirección de la empresa", "DireccionEmpresa", "Dirección")),
    (11, ("Correo empresarial", "Correo", "Email")),
    (12, ("Número de teléfono", "Telefono", "Teléfono")),
    (15, ("Moneda de cotización", "Moneda")),
    (16, ("Flete",)),
    (17, ("País de destino", "PaisDestino", "País de envío")),
    (18, ("Estado de destino", "EstadoDestino", "Estado de envío")),
    (19, ("Ciudad de destino", "CiudadDestino", "Ciudad de envío")),
    (22, ("Fianzas",)),
    (23, ("Póliza de responsabilidad civil", "PolizaResponsabilidadCivil")),
    (24, ("Monto de póliza", "MontoPoliza")),
    (25, ("Licitación", "Licitacion")),
    (26, ("Prioridad", "Prioridad de cotización")),
    (29, ("Fecha de entrega", "FechaEntrega")),
    (30, ("Fecha FDI", "Creación", "Created")),
    (33, ("Notas generales", "Notas", "Comentarios", "Consideraciones especiales")),
    (34, ("DocsUrl", "DocsLink", "Documentos")),
]

COMMON_LAYOUT = {
    "SEL": {"design_end": 63, "acabado": 66, "galv": 67, "tipo_galv": 68, "precio_galv": 69, "services": 78, "comments": 92},
    "DIN": {"design_end": 71, "acabado": 74, "galv": 75, "tipo_galv": 76, "precio_galv": 77, "services": 86, "comments": 100},
    "PBK": {"design_end": 71, "acabado": 74, "galv": 75, "tipo_galv": 76, "precio_galv": 77, "services": 86, "comments": 100},
    "DRV": {"design_end": 73, "acabado": 76, "galv": 77, "tipo_galv": 78, "precio_galv": 79, "services": 88, "comments": 102},
    "CAN": {"design_end": 62, "acabado": 65, "galv": 66, "tipo_galv": 67, "precio_galv": 68, "services": 77, "comments": 91},
    "MEZ": {"design_end": 74, "acabado": 77, "galv": 78, "tipo_galv": 79, "precio_galv": 80, "services": 89, "comments": 103},
    "CFL": {"design_end": 67, "acabado": 70, "galv": 71, "tipo_galv": 72, "precio_galv": 73, "services": 82, "comments": 96},
    "MZL": {"design_end": 86, "acabado": 89, "galv": 90, "tipo_galv": 91, "precio_galv": 92, "services": 101, "comments": 115},
}

def load_layout_config() -> dict[str, Any]:
    return json.loads(LAYOUT_FILE.read_text(encoding="utf-8"))


LAYOUT_CONFIG = load_layout_config()
TABLE_DEFINITIONS: dict[str, dict[str, Any]] = LAYOUT_CONFIG["tables"]


def truthy(value: Any) -> Any:
    if isinstance(value, bool):
        return "Sí" if value else "No"
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return value
    return str(value)


def is_truthy(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    return str(value).strip().lower() in ("sí", "si", "true", "1", "yes", "y")


def pick(data: dict[str, Any], *keys: str, default: Any = "") -> Any:
    for key in keys:
        if key in data and data[key] not in (None, ""):
            return data[key]
    return default


def as_list(value: Any) -> list[dict[str, Any]]:
    if isinstance(value, list):
        return [item if isinstance(item, dict) else {} for item in value]
    return []


def rows_from_summary(rows: list[dict[str, Any]], summary: Any) -> list[dict[str, Any]]:
    if rows:
        return rows
    text = str(summary or "").strip()
    if not text:
        return []
    result: list[dict[str, Any]] = []
    for line in text.splitlines():
        if not line.strip():
            continue
        parts = line.split(" - ")
        result.append({"Pieza": parts[0], "Comentarios": " - ".join(parts[1:])})
    return result


def html_to_text(value: Any) -> str:
    text = str(value or "")
    if not text:
        return ""
    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.IGNORECASE)
    text = re.sub(r"</t[dh]\s*>", " | ", text, flags=re.IGNORECASE)
    text = re.sub(r"</tr\s*>", "\n", text, flags=re.IGNORECASE)
    text = re.sub(r"</(p|div|li)\s*>", "\n", text, flags=re.IGNORECASE)
    text = re.sub(r"<li[^>]*>", "- ", text, flags=re.IGNORECASE)
    text = re.sub(r"<[^>]+>", "", text)
    text = unescape(text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    lines = [re.sub(r"\s*\|\s*$", "", line.strip()) for line in text.splitlines()]
    return "\n".join(line for line in lines if line).strip()


def normalize_method(system: dict[str, Any]) -> str:
    raw = str(pick(system, "TipoDiseño", "MetodoCaptura", "MetodoCotizacion", "MétodoCotización", default="Diseño")).strip().lower()
    if "pieza" in raw:
        return "ListaPiezas"
    if "pedido" in raw or "cotización anterior" in raw or "cotizacion anterior" in raw:
        return "PedidoAnterior"
    if "plano" in raw or "cliente" in raw:
        return "Planos"
    return "Diseño"


def method_label(system: dict[str, Any]) -> str:
    return {
        "ListaPiezas": "Listado de piezas",
        "PedidoAnterior": "Cotización o pedido anterior",
        "Planos": "Planos/diseño de cliente",
        "Diseño": "Diseño",
    }[normalize_method(system)]


def safe_sheet_name(name: str) -> str:
    cleaned = "".join("_" if ch in r'[]:*?/\\' else ch for ch in name)
    return cleaned[:31]


def copy_tables(source_ws, target_ws, suffix: str) -> None:
    for table in source_ws.tables.values():
        copied = deepcopy(table)
        table_name = f"{table.name}_{suffix}"
        copied.name = table_name
        copied.displayName = table_name
        target_ws.add_table(copied)


def table_columns(table_key: str) -> list[tuple[str, ...]]:
    return [tuple(aliases) for aliases in TABLE_DEFINITIONS[table_key]["aliases"]]


def table_headers(table_key: str) -> list[str]:
    return [str(header) for header in TABLE_DEFINITIONS[table_key]["headers"]]


def table_preview_headers(table_key: str) -> list[str]:
    definition = TABLE_DEFINITIONS[table_key]
    return [str(header) for header in definition.get("previewHeaders", definition["headers"])]


def find_table(ws, table_key: str, tipo: str):
    base_name = TABLE_DEFINITIONS[table_key]["baseName"]
    expected_prefix = f"{base_name}_{tipo}"
    for table in ws.tables.values():
        if table.name == expected_prefix or table.name.startswith(f"{expected_prefix}_"):
            return table
    for table in ws.tables.values():
        if table.name.startswith(f"{base_name}_"):
            return table
    return None


class RowInsertContext:
    def __init__(self) -> None:
        self.insertions: list[tuple[int, int]] = []

    def record(self, insert_at_row: int, count: int) -> None:
        self.insertions.append((insert_at_row, count))

    def adjust_ref(self, ref: str) -> str:
        min_col, min_row, max_col, max_row = range_boundaries(ref)
        for insert_at_row, count in self.insertions:
            if insert_at_row <= min_row:
                min_row += count
                max_row += count
            elif min_row < insert_at_row <= max_row + 1:
                max_row += count
        return format_range_ref(min_col, min_row, max_col, max_row)


def format_range_ref(min_col: int, min_row: int, max_col: int, max_row: int) -> str:
    return f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"


def set_cell(ws, address: str, value: Any) -> None:
    target = ws[address]
    if target.__class__.__name__ == "MergedCell":
        for merged_range in ws.merged_cells.ranges:
            if address in merged_range:
                target = ws.cell(merged_range.min_row, merged_range.min_col)
                break
    target.value = truthy(value)


def clear_table_area(ws, start_row: int, end_row: int, column_count: int) -> None:
    for row in range(start_row, end_row + 1):
        for col in range(1, column_count + 1):
            ws.cell(row, col).value = None


def clear_table_range(ws, ref: str, column_count: int) -> None:
    min_col, min_row, _max_col, max_row = range_boundaries(ref)
    for row in range(min_row + 1, max_row + 1):
        for col in range(min_col, min_col + column_count):
            ws.cell(row, col).value = None


def ensure_range_capacity(ws, ref: str, row_count: int, context: RowInsertContext) -> str:
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    capacity = max(0, max_row - min_row)
    extra_rows = max(0, row_count - capacity)
    if extra_rows:
        insert_at_row = max_row + 1
        ws.insert_rows(insert_at_row, amount=extra_rows)
        context.record(insert_at_row, extra_rows)
        max_row += extra_rows
    return format_range_ref(min_col, min_row, max_col, max_row)


def write_table_to_ref(ws, ref: str, rows: list[dict[str, Any]], columns: list[tuple[str, ...]], headers: list[str]) -> None:
    min_col, min_row, _max_col, max_row = range_boundaries(ref)
    capacity = max(0, max_row - min_row)
    if len(rows) > capacity:
        raise ValueError(f"El rango {ws.title}!{ref} no tiene filas suficientes para escribir {len(rows)} registros.")

    unmerge_intersecting_range(ws, ref)
    for col_offset, header in enumerate(headers, 1):
        ws.cell(min_row, min_col + col_offset - 1).value = header
    clear_table_range(ws, ref, len(columns))
    for row_offset, item in enumerate(rows):
        for col_offset, aliases in enumerate(columns, 1):
            ws.cell(min_row + 1 + row_offset, min_col + col_offset - 1).value = truthy(pick(item, *aliases))


def write_system_table_preview(ws, tipo: str, table_key: str, rows: list[dict[str, Any]], context: RowInsertContext) -> None:
    definition = TABLE_DEFINITIONS[table_key]
    preview_ref = definition.get("previewRanges", {}).get(tipo)
    if not preview_ref:
        return
    columns = table_columns(table_key)
    headers = table_preview_headers(table_key)
    ref = ensure_range_capacity(ws, context.adjust_ref(preview_ref), len(rows), context)
    write_table_to_ref(ws, ref, rows, columns, headers)


def write_system_table_data(ws, tipo: str, table_key: str, rows: list[dict[str, Any]], context: RowInsertContext) -> None:
    definition = TABLE_DEFINITIONS[table_key]
    columns = table_columns(table_key)
    headers = table_headers(table_key)
    table = find_table(ws, table_key, tipo)
    base_ref = table.ref if table is not None else definition["range"]
    ref = ensure_range_capacity(ws, context.adjust_ref(base_ref), len(rows), context)
    if table is not None:
        table.ref = ref
    write_table_to_ref(ws, ref, rows, columns, headers)


def set_rows_hidden(ws, start_row: int, end_row: int, hidden: bool) -> None:
    for row in range(start_row, end_row + 1):
        ws.row_dimensions[row].hidden = hidden


def apply_method_visibility(ws, tipo: str, system: dict[str, Any]) -> None:
    method = normalize_method(system)
    design_end = COMMON_LAYOUT[tipo]["design_end"]
    set_rows_hidden(ws, 11, 14, method != "PedidoAnterior")
    set_rows_hidden(ws, 15, 22, method != "ListaPiezas")
    set_rows_hidden(ws, 24, 26, method != "Planos")
    set_rows_hidden(ws, 27, design_end, method in ("PedidoAnterior", "ListaPiezas"))


def acabado_value(system: dict[str, Any]) -> str:
    raw = str(pick(system, "Acabado", "TipoGalv", default="")).strip()
    if raw:
        lower = raw.lower()
        if lower in ("frio", "frío", "galvanizado en frio", "galvanizado en frío"):
            return "Galvanizado en frío"
        if lower in ("caliente", "hot-dip", "hot dip", "galvanizado en caliente"):
            return "Galvanizado en caliente"
        if lower in ("pregalvanizado", "pre-galvanizado"):
            return "Pregalvanizado"
        if lower == "pintado":
            return "Pintado"
        return raw
    if is_truthy(pick(system, "GalvList", "Galvanizado", default=False)):
        return "Galvanizado"
    return "Pintado"


def galvanizado_from_acabado(acabado: str) -> bool:
    return "galv" in acabado.lower() and acabado.lower() != "pintado"


def tipo_galvanizado_from_acabado(acabado: str) -> str:
    lower = acabado.lower()
    if "fr" in lower:
        return "Frio"
    if "caliente" in lower or "hot" in lower:
        return "Caliente"
    if "pregalv" in lower:
        return "Pregalvanizado"
    return ""


def fill_header(wb, cotizacion: dict[str, Any]) -> None:
    ws = wb[HEADER_SHEET]
    for row, aliases in HEADER_ROWS:
        set_cell(ws, f"B{row}", pick(cotizacion, *aliases))


def fill_sistemas_index(wb, systems: list[dict[str, Any]], sheet_names: dict[int, str]) -> None:
    ws = wb[SYSTEMS_INDEX_SHEET]
    clear_table_area(ws, 6, max(ws.max_row, len(systems) + 5), 6)
    for index, system in enumerate(systems, 1):
        row = index + 5
        tipo = str(pick(system, "TipoKey", "tipo", "Tipo", default="SEL")).upper()
        ws.cell(row, 1).value = index
        ws.cell(row, 2).value = tipo
        ws.cell(row, 3).value = pick(system, "NombreSistema", "TipoNombre", "Title", default=f"Sistema {index}")
        ws.cell(row, 4).value = method_label(system)
        ws.cell(row, 5).value = sheet_names.get(index, "")
        ws.cell(row, 6).value = html_to_text(pick(system, "ConsEsp", "HTMLCol", "Comentarios", default=""))
    if "tblSistemas" in ws.tables:
        ws.tables["tblSistemas"].ref = f"A5:F{max(6, len(systems) + 5)}"


def fill_identity_and_method(ws, tipo: str, system: dict[str, Any], system_no: int) -> None:
    apply_method_visibility(ws, tipo, system)
    set_cell(ws, "B5", system_no)
    set_cell(ws, "B6", pick(system, "SistemaId", "SistemaID", "RegistroID", default=system_no))
    set_cell(ws, "B7", method_label(system))
    set_cell(ws, "B8", pick(system, "NombreSistema", "TipoNombre", "Title", default=f"Sistema {system_no}"))
    set_cell(ws, "B12", pick(system, "FolioCotizacionAnterior", "FolioCotAnterior"))
    set_cell(ws, "B13", pick(system, "FolioPedidoAnterior", "PedidoBase", "NumPedidoCot"))
    set_cell(ws, "B14", pick(system, "ComentariosReferencia", "ConsEsp"))
    set_cell(ws, "B25", truthy(pick(system, "AdjuntarImagenLayout", "RequiereAdjuntarLayout")))
    set_cell(ws, "B26", pick(system, "DocsUrl", "DocsLink", "Documentos"))


def fill_common(ws, tipo: str, system: dict[str, Any]) -> None:
    layout = COMMON_LAYOUT[tipo]
    acabado = acabado_value(system)
    tipo_galv = tipo_galvanizado_from_acabado(acabado)
    set_cell(ws, f"B{layout['acabado']}", acabado)
    set_cell(ws, f"B{layout['galv']}", galvanizado_from_acabado(acabado))
    set_cell(ws, f"B{layout['tipo_galv']}", tipo_galv)
    set_cell(ws, f"B{layout['precio_galv']}", pick(system, "PpkgGalv", "Precio por kilogramo galvanizado") if tipo_galv in ("Frio", "Caliente") else "")
    service_row = layout["services"]
    set_cell(ws, f"B{service_row}", pick(system, "Ins", "Instalacion", "Instalación"))
    set_cell(ws, f"B{service_row + 1}", pick(system, "CostoInstalacion", "Costo instalación"))
    set_cell(ws, f"B{service_row + 2}", pick(system, "ComentariosInstalacion", "Comentarios instalación"))
    set_cell(ws, f"B{service_row + 3}", pick(system, "MemCalc", "Memoria de cálculo"))
    set_cell(ws, f"B{service_row + 4}", pick(system, "CostoMemCalculo", "Costo memoria cálculo"))
    set_cell(ws, f"B{service_row + 5}", pick(system, "EstProv", "Unirse a estructura de otro proveedor"))
    set_cell(ws, f"B{service_row + 6}", pick(system, "ComentariosEstructura", "Comentarios estructura"))
    set_cell(ws, f"B{service_row + 7}", pick(system, "ProvExternos", "Proveedores externos"))
    set_cell(ws, f"B{layout['comments']}", pick(system, "ConsEsp", "Consideraciones especiales", "Comentarios"))


def fill_area_and_levels(ws, system: dict[str, Any], area_row: int, levels_row: int, include_montacargas: bool) -> None:
    if include_montacargas:
        set_cell(ws, f"B{area_row}", pick(system, "PasilloMax", "Pasillo máximo"))
        set_cell(ws, f"B{area_row + 1}", pick(system, "PasilloMin", "Pasillo mínimo"))
        set_cell(ws, f"B{area_row + 2}", pick(system, "AnchoDisp", "Ancho disponible"))
        set_cell(ws, f"B{area_row + 3}", pick(system, "LargoDisp", "Largo disponible"))
        set_cell(ws, f"B{levels_row}", pick(system, "ConsiderarAlturaMaxMonta"))
        set_cell(ws, f"B{levels_row + 1}", pick(system, "AltCritMonta", "Altura crítica de montacargas"))
        offset = 2
    else:
        set_cell(ws, f"B{area_row}", pick(system, "AnchoDisp", "Ancho disponible"))
        set_cell(ws, f"B{area_row + 1}", pick(system, "LargoDisp", "Largo disponible"))
        set_cell(ws, f"B{area_row + 2}", pick(system, "AnchoPasilloPickeo", "Ancho pasillo pickeo"))
        offset = 0
    set_cell(ws, f"B{levels_row + offset}", pick(system, "ConsiderarAlturaNave"))
    set_cell(ws, f"B{levels_row + offset + 1}", pick(system, "AlturaMaxNave", "Altura máxima nave"))
    set_cell(ws, f"B{levels_row + offset + 2}", pick(system, "AlturaMinNave", "Altura mínima nave"))
    set_cell(ws, f"B{levels_row + offset + 3}", pick(system, "AdjuntarImagenLayout", "Requiere adjuntar layout"))
    set_cell(ws, f"B{levels_row + offset + 4}", pick(system, "ExisteDefCliente", "Existe definición cliente"))
    set_cell(ws, f"B{levels_row + offset + 5}", pick(system, "ComentariosConfigCliente", "DefPorCliente", "Definido por el cliente"))


def fill_rack(ws, system: dict[str, Any], start_row: int, high_impact: bool = True) -> None:
    set_cell(ws, f"B{start_row}", pick(system, "FrentesBuscados", "Frentes buscados"))
    set_cell(ws, f"B{start_row + 1}", pick(system, "FondosBuscados", "Fondos buscados"))
    set_cell(ws, f"B{start_row + 2}", pick(system, "NivelesBuscados", "Niveles buscados"))
    set_cell(ws, f"B{start_row + 3}", pick(system, "TipoRodamiento", "Tipo rodamiento"))
    set_cell(ws, f"B{start_row + 4}", pick(system, "MetodoCalculoEntrecentros", "Método cálculo entrecentros"))
    if high_impact:
        set_cell(ws, f"B{start_row + 5}", pick(system, "UtilizarRodamientoAltoImpacto"))
        set_cell(ws, f"B{start_row + 6}", pick(system, "EspecificacionRodamientoAltoImpacto", "Especificación rodamiento alto impacto"))


def fill_detail_tables(ws, tipo: str, system: dict[str, Any]) -> None:
    context = RowInsertContext()
    table_writes = [
        (
            "piezas",
            rows_from_summary(as_list(pick(system, "piezas", "Piezas", default=[])), pick(system, "PiezasResumen")),
        ),
        ("tarimas", as_list(pick(system, "tarimas", "Tarimas", default=[]))),
        ("productos", as_list(pick(system, "productos", "Productos", default=[]))),
        (
            "elementosSeguridad",
            rows_from_summary(
                as_list(pick(system, "elementosSeguridad", "ElementosSeguridad", default=[])),
                pick(system, "ElementosSeguridadResumen"),
            ),
        ),
        ("piezasEspeciales", as_list(pick(system, "piezasEspeciales", "PiezasEspeciales", default=[]))),
        ("colores", as_list(pick(system, "colores", "Colores", default=[]))),
        (
            "proveedoresExternos",
            as_list(pick(system, "proveedoresExternos", "ProveedoresExternos", default=[])),
        ),
    ]
    for table_key, rows in table_writes:
        write_system_table_preview(ws, tipo, table_key, rows, context)
    for table_key, rows in table_writes:
        write_system_table_data(ws, tipo, table_key, rows, context)


def fill_supported_form(ws, tipo: str, system: dict[str, Any], system_no: int) -> None:
    fill_identity_and_method(ws, tipo, system, system_no)
    if tipo == "SEL":
        fill_area_and_levels(ws, system, 38, 43, True)
    elif tipo in ("DIN", "PBK"):
        fill_rack(ws, system, 38, True)
        fill_area_and_levels(ws, system, 46, 51, True)
    elif tipo == "DRV":
        set_cell(ws, "B38", pick(system, "FrentesBuscados", "Frentes buscados"))
        set_cell(ws, "B39", pick(system, "FondosBuscados", "Fondos buscados"))
        set_cell(ws, "B40", pick(system, "NivelesBuscados", "Niveles buscados"))
        set_cell(ws, "B42", pick(system, "TipoCapturaMontacargas"))
        set_cell(ws, "B43", pick(system, "AlturaCabinaMontacargas"))
        set_cell(ws, "B44", pick(system, "AnchoTotalMontacargas"))
        set_cell(ws, "B45", pick(system, "AnchoMastilMontacargas"))
        set_cell(ws, "B46", pick(system, "ModeloMontacargas"))
        fill_area_and_levels(ws, system, 48, 53, True)
    elif tipo == "CAN":
        set_cell(ws, "B29", pick(system, "TipoProducto"))
        set_cell(ws, "B30", pick(system, "LongitudCarga"))
        set_cell(ws, "B31", pick(system, "SeccionCarga"))
        set_cell(ws, "B32", pick(system, "PesoCarga"))
        set_cell(ws, "B33", pick(system, "CantidadPorNivel"))
        set_cell(ws, "B35", pick(system, "TipoGondola", "Tipo góndola"))
        set_cell(ws, "B37", pick(system, "PasilloMax", "Pasillo máximo"))
        set_cell(ws, "B38", pick(system, "PasilloMin", "Pasillo mínimo"))
        set_cell(ws, "B39", pick(system, "AnchoDisp", "Ancho disponible"))
        set_cell(ws, "B40", pick(system, "LargoDisp", "Largo disponible"))
        set_cell(ws, "B42", pick(system, "ConsiderarAlturaMaxMonta"))
        set_cell(ws, "B43", pick(system, "AltCritMonta", "Altura crítica de montacargas"))
        set_cell(ws, "B44", pick(system, "ConsiderarAlturaNave"))
        set_cell(ws, "B45", pick(system, "AlturaMaxNave", "Altura máxima nave"))
        set_cell(ws, "B46", pick(system, "AlturaMinNave", "Altura mínima nave"))
        set_cell(ws, "B47", pick(system, "AdjuntarImagenLayout", "Requiere adjuntar layout"))
        set_cell(ws, "B48", pick(system, "ExisteDefCliente", "Existe definición cliente"))
        set_cell(ws, "B49", pick(system, "ComentariosConfigCliente", "DefPorCliente", "Definido por el cliente"))
    elif tipo == "MEZ":
        set_cell(ws, "B39", pick(system, "AlturaRecomendadaEntrepiso", "Altura recomendada entrepiso"))
        set_cell(ws, "B40", pick(system, "CantidadEntrepisos", "Cantidad entrepisos"))
        set_cell(ws, "B41", pick(system, "RequiereElevador"))
        set_cell(ws, "B42", pick(system, "EspecificacionElevador", "Especificación elevador"))
        set_cell(ws, "B43", pick(system, "TipoPiso", "Tipo piso"))
        set_cell(ws, "B44", pick(system, "UsaCarrito"))
        set_cell(ws, "B45", pick(system, "MedidasCarrito"))
        set_cell(ws, "B46", pick(system, "NumeroRuedas", "Número ruedas"))
        set_cell(ws, "B47", pick(system, "TipoRueda"))
        set_cell(ws, "B48", pick(system, "MedidaRueda"))
        set_cell(ws, "B49", pick(system, "PesoCarrito"))
        set_cell(ws, "B50", pick(system, "RequiereEscaleras"))
        fill_area_and_levels(ws, system, 52, 56, False)
    elif tipo == "CFL":
        fill_rack(ws, system, 39, False)
        fill_area_and_levels(ws, system, 45, 49, False)
    elif tipo == "MZL":
        set_cell(ws, "B29", pick(system, "CantidadPisos"))
        set_cell(ws, "B30", pick(system, "CargaPorM2"))
        set_cell(ws, "B31", pick(system, "EsModulado"))
        set_cell(ws, "B32", pick(system, "ZonaModulada"))
        set_cell(ws, "B33", pick(system, "UsaTarimas"))
        set_cell(ws, "B34", pick(system, "RequiereElevador"))
        set_cell(ws, "B35", pick(system, "EspecificacionElevador", "Especificación elevador"))
        set_cell(ws, "B36", pick(system, "TipoPiso", "Tipo piso"))
        set_cell(ws, "B37", pick(system, "UsaCarrito"))
        set_cell(ws, "B38", pick(system, "MedidasCarrito"))
        set_cell(ws, "B39", pick(system, "NumeroRuedas", "Número ruedas"))
        set_cell(ws, "B40", pick(system, "TipoRueda"))
        set_cell(ws, "B41", pick(system, "MedidaRueda"))
        set_cell(ws, "B42", pick(system, "PesoCarrito"))
        set_cell(ws, "B43", pick(system, "RequiereEscaleras"))
        set_cell(ws, "B45", pick(system, "MetodoSeparacionColumnas", "Método separación columnas"))
        set_cell(ws, "B46", pick(system, "SeparacionColumnasManual", "Separación columnas manual"))
        fill_area_and_levels(ws, system, 64, 68, False)
    fill_common(ws, tipo, system)
    fill_detail_tables(ws, tipo, system)


def add_ot_sheet(wb, system: dict[str, Any], sheet_name: str, system_no: int) -> None:
    ws = wb.create_sheet(safe_sheet_name(sheet_name))
    ws["A1"] = "Formulario de Sistema (Otro)"
    ws["A3"] = "Sistema No"
    ws["B3"] = system_no
    ws["A4"] = "SistemaID"
    ws["B4"] = pick(system, "SistemaId", "SistemaID", "RegistroID")
    ws["A5"] = "Descripción"
    ws["B5"] = pick(system, "NombreSistema", "Title")
    ws["A6"] = "Documentos"
    ws["B6"] = pick(system, "DocsUrl", "DocsLink", "Documentos")
    ws["A8"] = "Detalle"
    ws["B8"] = html_to_text(pick(system, "HTMLCol", "HTML"))
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 80


def add_payload_sheet(wb, payload: dict[str, Any]) -> None:
    if PAYLOAD_SHEET in wb.sheetnames:
        del wb[PAYLOAD_SHEET]
    ws = wb.create_sheet(PAYLOAD_SHEET)
    ws.sheet_state = "hidden"
    text = json.dumps(payload, ensure_ascii=False, indent=2)
    ws["A1"] = "Payload JSON usado para generar este archivo"
    for index in range(0, len(text), 30000):
        ws.cell(index // 30000 + 2, 1).value = text[index:index + 30000]


def hide_layout_columns(ws) -> None:
    for column_range in LAYOUT_CONFIG.get("hiddenColumns", []):
        start, end = column_range.split(":", 1)
        for column_index in range(column_index_from_string(start), column_index_from_string(end) + 1):
            column_letter = get_column_letter(column_index)
            ws.column_dimensions[column_letter].hidden = True


def unmerge_intersecting_range(ws, ref: str) -> None:
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    for merged_range in list(ws.merged_cells.ranges):
        if (
            merged_range.max_col < min_col
            or merged_range.min_col > max_col
            or merged_range.max_row < min_row
            or merged_range.min_row > max_row
        ):
            continue
        try:
            ws.unmerge_cells(str(merged_range))
        except KeyError:
            if merged_range in ws.merged_cells.ranges:
                ws.merged_cells.ranges.remove(merged_range)
            for row in range(merged_range.min_row, merged_range.max_row + 1):
                for col in range(merged_range.min_col, merged_range.max_col + 1):
                    cell = ws._cells.get((row, col))
                    if isinstance(cell, MergedCell):
                        del ws._cells[(row, col)]


def ensure_preview_ranges(ws, tipo: str) -> None:
    for table_key, definition in TABLE_DEFINITIONS.items():
        ref = definition.get("previewRanges", {}).get(tipo)
        if not ref:
            continue
        min_col, min_row, _max_col, max_row = range_boundaries(ref)
        headers = table_preview_headers(table_key)
        unmerge_intersecting_range(ws, ref)
        for offset, header in enumerate(headers):
            ws.cell(min_row, min_col + offset).value = header
        for row in range(min_row + 1, max_row + 1):
            for col in range(min_col, min_col + len(headers)):
                ws.cell(row, col).value = None


def ensure_form_tables(wb) -> None:
    style = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    for tipo, sheet_name in TYPE_TEMPLATES.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        for table_name in list(ws.tables.keys()):
            del ws.tables[table_name]
        hide_layout_columns(ws)

        for table_key, definition in TABLE_DEFINITIONS.items():
            ref = definition["range"]
            min_col, min_row, _max_col, max_row = range_boundaries(ref)
            headers = definition["headers"]
            for offset, header in enumerate(headers):
                ws.cell(min_row, min_col + offset).value = header
            for row in range(min_row + 1, max_row + 1):
                for col in range(min_col, min_col + len(headers)):
                    ws.cell(row, col).value = None

            table_name = f"{definition['baseName']}_{tipo}"
            table = Table(displayName=table_name, ref=ref)
            table.tableStyleInfo = deepcopy(style)
            ws.add_table(table)
        ensure_preview_ranges(ws, tipo)


def write_layout_config_sheet(wb) -> None:
    config_sheet = LAYOUT_CONFIG.get("configSheet", "FDI_Table_Config")
    if config_sheet in wb.sheetnames:
        del wb[config_sheet]
    ws = wb.create_sheet(config_sheet)
    ws.sheet_state = "hidden"
    ws["A1"] = "FDI Excel layout JSON"
    text = json.dumps(LAYOUT_CONFIG, ensure_ascii=False, separators=(",", ":"))
    for index in range(0, len(text), 30000):
        ws.cell(index // 30000 + 2, 1).value = text[index:index + 30000]


def normalize_template(template: Path) -> None:
    wb = load_workbook(template)
    ensure_form_tables(wb)
    write_layout_config_sheet(wb)
    wb.save(template)


def create_system_sheets(wb, systems: list[dict[str, Any]]) -> dict[int, str]:
    sheet_names: dict[int, str] = {}
    template_for_type: dict[str, Any] = {}
    used_types: set[str] = set()
    to_fill: list[tuple[Any, str, dict[str, Any], int]] = []
    for index, system in enumerate(systems, 1):
        tipo = str(pick(system, "TipoKey", "tipo", "Tipo", default="SEL")).upper()
        sheet_name = safe_sheet_name(f"{tipo}_S{index:02d}_Form")
        sheet_names[index] = sheet_name
        if tipo in TYPE_TEMPLATES:
            template_name = TYPE_TEMPLATES[tipo]
            if tipo not in template_for_type:
                ws = wb[template_name]
                ws.title = sheet_name
                template_for_type[tipo] = ws
                used_types.add(tipo)
            else:
                ws = wb.copy_worksheet(template_for_type[tipo])
                ws.title = sheet_name
                copy_tables(template_for_type[tipo], ws, f"{index:02d}")
            to_fill.append((ws, tipo, system, index))
        else:
            add_ot_sheet(wb, system, sheet_name, index)
    for tipo, template_name in TYPE_TEMPLATES.items():
        if tipo not in used_types and template_name in wb.sheetnames:
            del wb[template_name]
    for ws, tipo, system, index in to_fill:
        fill_supported_form(ws, tipo, system, index)
    return sheet_names


def generate(template: Path, payload_path: Path, output: Path) -> None:
    payload = json.loads(payload_path.read_text(encoding="utf-8"))
    cotizacion = payload.get("cotizacion", {})
    systems = payload.get("sistemas", [])

    wb = load_workbook(template)
    ensure_form_tables(wb)
    write_layout_config_sheet(wb)
    fill_header(wb, cotizacion)
    sheet_names = create_system_sheets(wb, systems)
    fill_sistemas_index(wb, systems, sheet_names)
    add_payload_sheet(wb, payload)

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description="Genera una copia FDI XLSX desde la plantilla y un payload JSON.")
    parser.add_argument("--template", default="templates/FDI_Master.xlsx", type=Path)
    parser.add_argument("--payload", default="examples/fdi-payload.sample.json", type=Path)
    parser.add_argument("--output", default="work/fdi_closure_20260609_01/output/FDI_TEST_001-26.xlsx", type=Path)
    parser.add_argument("--normalize-template", action="store_true", help="Actualiza la plantilla con las tablas normalizadas y la configuración oculta.")
    args = parser.parse_args()
    if args.normalize_template:
        normalize_template(args.template)
        print(args.template)
        return
    generate(args.template, args.payload, args.output)
    print(args.output)


if __name__ == "__main__":
    main()
